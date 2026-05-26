import os
import sys
import json
import csv
import argparse
import zipfile
import xml.etree.ElementTree as ET
import asyncio
import subprocess
import logging

# Setup logger
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger("artifact-pipeline")

# Try importing docx and openpyxl, setup fallback flags
try:
    import docx

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import websockets

    HAS_WEBSOCKETS = True
except ImportError:
    HAS_WEBSOCKETS = False


def extract_docx_metadata(file_path: str) -> dict:
    """Fallback docx parser that extracts metadata and raw xml text from the zip archive directly."""
    logger.info(f"Using zip-based fallback metadata extractor for docx: {file_path}")
    metadata = {}
    try:
        with zipfile.ZipFile(file_path) as zf:
            # Parse core properties XML
            if "docProps/core.xml" in zf.namelist():
                core_xml = zf.read("docProps/core.xml")
                root = ET.fromstring(core_xml)
                ns = {
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "dcterms": "http://purl.org/dc/terms/",
                }
                metadata["creator"] = (
                    root.find(".//dc:creator", ns).text
                    if root.find(".//dc:creator", ns) is not None
                    else "Unknown"
                )
                metadata["created"] = (
                    root.find(".//dcterms:created", ns).text
                    if root.find(".//dcterms:created", ns) is not None
                    else "Unknown"
                )
                metadata["modified"] = (
                    root.find(".//dcterms:modified", ns).text
                    if root.find(".//dcterms:modified", ns) is not None
                    else "Unknown"
                )
                metadata["title"] = (
                    root.find(".//dc:title", ns).text
                    if root.find(".//dc:title", ns) is not None
                    else "Untitled"
                )

            # Simple word/document.xml extraction
            if "word/document.xml" in zf.namelist():
                doc_xml = zf.read("word/document.xml")
                root = ET.fromstring(doc_xml)
                # Extract paragraph text elements
                paragraphs = []
                for p in root.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                ):
                    if p.text:
                        paragraphs.append(p.text)
                metadata["raw_text_sample"] = " ".join(paragraphs[:50]) + (
                    "..." if len(paragraphs) > 50 else ""
                )
                metadata["parsing_status"] = "Fallback (Metadata Only)"
        return metadata
    except Exception as e:
        logger.error(f"Zip metadata extraction failed: {e}")
        return {
            "error": f"Metadata extraction failed: {str(e)}",
            "parsing_status": "Failed",
        }


def extract_xlsx_metadata(file_path: str) -> dict:
    """Fallback xlsx parser that extracts metadata from the workbook zip package."""
    logger.info(f"Using zip-based fallback metadata extractor for xlsx: {file_path}")
    metadata = {}
    try:
        with zipfile.ZipFile(file_path) as zf:
            if "docProps/core.xml" in zf.namelist():
                core_xml = zf.read("docProps/core.xml")
                root = ET.fromstring(core_xml)
                ns = {
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "dcterms": "http://purl.org/dc/terms/",
                }
                metadata["creator"] = (
                    root.find(".//dc:creator", ns).text
                    if root.find(".//dc:creator", ns) is not None
                    else "Unknown"
                )
                metadata["created"] = (
                    root.find(".//dcterms:created", ns).text
                    if root.find(".//dcterms:created", ns) is not None
                    else "Unknown"
                )
                metadata["modified"] = (
                    root.find(".//dcterms:modified", ns).text
                    if root.find(".//dcterms:modified", ns) is not None
                    else "Unknown"
                )
                metadata["parsing_status"] = "Fallback (Metadata Only)"
        return metadata
    except Exception as e:
        logger.error(f"Xlsx zip metadata extraction failed: {e}")
        return {
            "error": f"Metadata extraction failed: {str(e)}",
            "parsing_status": "Failed",
        }


def parse_docx(file_path: str) -> dict:
    """Parses DOCX requirements template."""
    if not HAS_DOCX:
        logger.warning(
            "python-docx library not found. Falling back to metadata extraction."
        )
        return extract_docx_metadata(file_path)

    try:
        doc = docx.Document(file_path)
        requirements = []
        current_req = {}

        for p in doc.paragraphs:
            text = p.text.strip()
            if text.startswith("SYS_REQ_") or text.startswith("REQ_"):
                if current_req:
                    requirements.append(current_req)
                parts = text.split(":", 1)
                req_id = parts[0].strip()
                req_title = parts[1].strip() if len(parts) > 1 else ""
                current_req = {
                    "id": req_id,
                    "title": req_title,
                    "description": "",
                    "aspice_ref": "SYS.2",
                }
            elif current_req and text:
                if text.startswith("ASPICE Reference:"):
                    current_req["aspice_ref"] = text.replace(
                        "ASPICE Reference:", ""
                    ).strip()
                else:
                    current_req["description"] += text + "\n"

        if current_req:
            requirements.append(current_req)

        return {
            "parsing_status": "Success",
            "type": "docx",
            "requirements": requirements,
        }
    except Exception as e:
        logger.error(f"Error parsing docx using python-docx: {e}. Attempting fallback.")
        return extract_docx_metadata(file_path)


def parse_xlsx(file_path: str) -> dict:
    """Parses XLSX requirements sheet."""
    if not HAS_OPENPYXL:
        logger.warning(
            "openpyxl library not found. Falling back to metadata extraction."
        )
        return extract_xlsx_metadata(file_path)

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        requirements = []
        # Expect first sheet
        ws = wb.active
        headers = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx == 0:
                headers = [str(cell).lower() for cell in row]
                continue

            row_dict = {}
            for c_idx, cell in enumerate(row):
                if c_idx < len(headers):
                    row_dict[headers[c_idx]] = cell

            if row_dict.get("id") or row_dict.get("requirement id"):
                requirements.append(
                    {
                        "id": row_dict.get("id") or row_dict.get("requirement id"),
                        "title": row_dict.get("title") or row_dict.get("name") or "",
                        "description": row_dict.get("description")
                        or row_dict.get("text")
                        or "",
                        "aspice_ref": row_dict.get("aspice_ref")
                        or row_dict.get("aspice")
                        or "SYS.2",
                    }
                )
        return {
            "parsing_status": "Success",
            "type": "xlsx",
            "requirements": requirements,
        }
    except Exception as e:
        logger.error(f"Error parsing xlsx: {e}. Attempting fallback.")
        return extract_xlsx_metadata(file_path)


def parse_csv(file_path: str) -> dict:
    """Parses CSV requirements sheet."""
    try:
        requirements = []
        with open(file_path, mode="r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                requirements.append(
                    {
                        "id": row.get("id") or row.get("Requirement ID") or "",
                        "title": row.get("title") or row.get("Name") or "",
                        "description": row.get("description")
                        or row.get("Description")
                        or "",
                        "aspice_ref": row.get("aspice_ref")
                        or row.get("ASPICE Reference")
                        or "SYS.2",
                    }
                )
        return {
            "parsing_status": "Success",
            "type": "csv",
            "requirements": requirements,
        }
    except Exception as e:
        logger.error(f"Error parsing csv: {e}")
        return {"error": f"CSV parse failed: {str(e)}", "parsing_status": "Failed"}


def parse_artifact(file_path: str) -> dict:
    """Dispatches parsing based on file extension."""
    if not os.path.exists(file_path):
        return {"error": f"File {file_path} not found.", "parsing_status": "Failed"}

    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".xlsx":
        return parse_xlsx(file_path)
    elif ext == ".csv":
        return parse_csv(file_path)
    else:
        return {"error": f"Unsupported file type: {ext}", "parsing_status": "Failed"}


def run_validation(repo_path: str) -> dict:
    """Runs verification checks: Ruff linting and pytest test execution."""
    logger.info("Running validation pipeline...")
    results = {"lint": "unknown", "tests": "unknown", "errors": []}

    # Linting check
    try:
        venv_python = os.path.join(repo_path, ".venv", "Scripts", "python.exe")
        if not os.path.exists(venv_python):
            venv_python = "python"

        lint_cmd = [venv_python, "-m", "ruff", "check", "."]
        lint_res = subprocess.run(
            lint_cmd, cwd=repo_path, capture_output=True, text=True
        )
        results["lint"] = "pass" if lint_res.returncode == 0 else "fail"
        if lint_res.returncode != 0:
            results["errors"].append(f"Linting failed: {lint_res.stdout}")
    except Exception as e:
        results["lint"] = "error"
        results["errors"].append(f"Lint execution error: {str(e)}")

    # Testing check
    try:
        pytest_cmd = [venv_python, "-m", "pytest", "--cov=runtime"]
        test_res = subprocess.run(
            pytest_cmd, cwd=repo_path, capture_output=True, text=True
        )
        results["tests"] = "pass" if test_res.returncode == 0 else "fail"
        if test_res.returncode != 0:
            results["errors"].append(f"Tests failed: {test_res.stdout}")
    except Exception as e:
        results["tests"] = "error"
        results["errors"].append(f"Test execution error: {str(e)}")

    return results


# WebSocket broadcasting server logic
async def ws_handler(websocket, path):
    logger.info("WebUI Client connected to websocket gateway")
    # Send initial state
    await websocket.send(
        json.dumps(
            {
                "type": "log",
                "message": "Connected to ASPICE systems lead pipeline gateway",
            }
        )
    )

    try:
        async for message in websocket:
            data = json.loads(message)
            cmd = data.get("command")
            logger.info(f"WebSocket command received: {cmd}")

            if cmd == "parse":
                file_path = data.get("file")
                await websocket.send(
                    json.dumps({"type": "status", "status": "parsing"})
                )
                await websocket.send(
                    json.dumps({"type": "log", "message": f"Parsing file: {file_path}"})
                )
                result = parse_artifact(file_path)
                await websocket.send(
                    json.dumps({"type": "parsing_result", "result": result})
                )
                await websocket.send(json.dumps({"type": "status", "status": "idle"}))

            elif cmd == "validate":
                await websocket.send(
                    json.dumps({"type": "status", "status": "validating"})
                )
                await websocket.send(
                    json.dumps(
                        {
                            "type": "log",
                            "message": "Starting verification (ruff check + pytest)...",
                        }
                    )
                )
                repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
                res = run_validation(repo_root)
                await websocket.send(
                    json.dumps({"type": "validation_result", "result": res})
                )
                await websocket.send(json.dumps({"type": "status", "status": "idle"}))

            else:
                await websocket.send(
                    json.dumps({"type": "error", "message": f"Unknown command: {cmd}"})
                )
    except Exception as e:
        logger.error(f"WebSocket connection error: {e}")
    finally:
        logger.info("WebUI Client disconnected")


async def main_async(port):
    if not HAS_WEBSOCKETS:
        logger.error("websockets library not installed. Cannot start websocket server.")
        sys.exit(1)

    logger.info(f"Starting websocket server on port {port}...")
    async with websockets.serve(ws_handler, "localhost", port):
        await asyncio.Future()  # run forever


def main():
    parser = argparse.ArgumentParser(
        description="ASPICE Technical Systems Lead Artifact Pipeline"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    # Subcommand: parse
    p_parse = subparsers.add_parser("parse", help="Parse requirements artifact")
    p_parse.add_argument(
        "--file", required=True, help="Path to docx/xlsx/csv artifact file"
    )
    p_parse.add_argument("--output", required=True, help="Output path for JSON result")

    # Subcommand: validate
    p_val = subparsers.add_parser(
        "validate", help="Run linting and test validation checks"
    )
    p_val.add_argument(
        "--repo-path", required=True, help="Path to agent-shell workspace root"
    )
    p_val.add_argument(
        "--output", required=True, help="Output path for JSON validation result"
    )

    # Subcommand: serve-websocket
    p_serve = subparsers.add_parser(
        "serve-websocket", help="Start the WebSocket server gateway"
    )
    p_serve.add_argument(
        "--port", type=int, default=8765, help="Port to run WebSocket server on"
    )

    args = parser.parse_args()

    if args.command == "parse":
        res = parse_artifact(args.file)
        with open(args.output, "w", encoding="utf-8") as f:
            json.dump(res, f, indent=2)
        print(f"Parsing complete. Results written to {args.output}")

    elif args.command == "validate":
        res = run_validation(args.repo_path)
        with open(args.output, "w", encoding="utf-8") as f:
            json.dump(res, f, indent=2)
        print(f"Validation complete. Results written to {args.output}")

    elif args.command == "serve-websocket":
        asyncio.run(main_async(args.port))


if __name__ == "__main__":
    main()
